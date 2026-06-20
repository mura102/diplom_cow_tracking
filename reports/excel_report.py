# reports/excel_report.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime


COLOR_HEADER   = "1F4E79"
COLOR_ROW_ODD  = "EBF3FB"
COLOR_ROW_EVEN = "FFFFFF"

THIN_BORDER = Border(
    left=Side(style="thin", color="B0C4DE"),
    right=Side(style="thin", color="B0C4DE"),
    top=Side(style="thin", color="B0C4DE"),
    bottom=Side(style="thin", color="B0C4DE"),
)

ACTIVITY_MAP = {
    "FEED":  "Кормление",
    "DRINK": "Водопой",
    "SLEEP": "Сон/Отдых",
    "FALL":  "Падение",
    "BCS":   "Оценка BCS",
}


def _header(cell, text: str):
    cell.value = text
    cell.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    cell.fill = PatternFill("solid", fgColor=COLOR_HEADER)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER


def _data(cell, value, row_idx: int):
    cell.value = value
    if row_idx % 2 == 0:
        cell.fill = PatternFill("solid", fgColor=COLOR_ROW_ODD)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.font = Font(name="Calibri", size=10)
    cell.border = THIN_BORDER


def _write_table(ws, start_row: int, headers: list[str], widths: list[int], rows: list[list]) -> int:
    """Пишет таблицу начиная с start_row, возвращает следующую свободную строку."""
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=start_row, column=col)
        _header(c, h)
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[start_row].height = 20

    if not rows:
        last_col = len(headers)
        ws.merge_cells(start_row=start_row + 1, start_column=1, end_row=start_row + 1, end_column=last_col)
        c = ws.cell(row=start_row + 1, column=1)
        c.value = "Нет данных за выбранный период"
        c.alignment = Alignment(horizontal="center")
        c.font = Font(italic=True, color="999999")
        return start_row + 3

    for idx, row_values in enumerate(rows, start=1):
        row = start_row + idx
        for col, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row, column=col)
            _data(cell, val, idx)
        ws.row_dimensions[row].height = 18

    return start_row + len(rows) + 2


def generate_excel_report(
    activity_stats: list[dict],
    cow_summary: list[dict],
    save_path: str,
    company_name: str = "УМНАЯ ФЕРМА",
    period_days: int = 7,
) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    ws.sheet_view.showGridLines = False

    # ── Заголовок ──
    ws.merge_cells("A1:F1")
    t = ws["A1"]
    t.value = f"🐄 {company_name}"
    t.font = Font(bold=True, size=14, color="1F4E79", name="Calibri")
    t.alignment = Alignment(horizontal="center", vertical="center")
    t.fill = PatternFill("solid", fgColor="D6E4F0")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:F2")
    d = ws["A2"]
    d.value = (
        f"Отчёт по статистике активности животных | "
        f"Период: последние {period_days} дней | "
        f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    )
    d.font = Font(italic=True, size=9, color="666666", name="Calibri")
    d.alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 18

    ws.row_dimensions[3].height = 5

    # ── 1. Сводка по стаду ──
    ws.merge_cells("A4:F4")
    s1 = ws["A4"]
    s1.value = "1. Сводка по стаду"
    s1.font = Font(bold=True, size=12, color="1F4E79", name="Calibri")
    ws.row_dimensions[4].height = 22

    headers1 = ["ID", "Порода", "Вес, кг", "Статус", "Последний осмотр", "Событий"]
    widths1  = [8, 18, 10, 16, 20, 10]
    rows1 = []
    for cow in cow_summary:
        dt = cow.get("last_inspection")
        dt_str = dt.strftime("%d.%m.%Y %H:%M") if hasattr(dt, "strftime") else (str(dt) if dt else "—")
        weight = cow.get("weight_kg")
        weight_str = f"{weight:.0f}" if isinstance(weight, (int, float)) else "—"
        rows1.append([
            cow.get("cow_number", "—"),
            cow.get("breed", "—"),
            weight_str,
            cow.get("status", "—"),
            dt_str,
            cow.get("total_events", 0),
        ])
    next_row = _write_table(ws, 5, headers1, widths1, rows1)

    # ── 2. Детальная статистика активности (справа от первой таблицы) ──
    second_table_col = 8  # столбец H

    ws.merge_cells(start_row=4, start_column=second_table_col, end_row=4, end_column=second_table_col + 3)
    s2 = ws.cell(row=4, column=second_table_col)
    s2.value = "2. Детальная статистика активности"
    s2.font = Font(bold=True, size=12, color="1F4E79", name="Calibri")

    headers2 = ["ID", "Дата", "Тип активности", "Кол-во событий"]
    widths2 = [8, 14, 22, 16]

    for col_offset, (h, w) in enumerate(zip(headers2, widths2)):
        col = second_table_col + col_offset
        c = ws.cell(row=5, column=col)
        _header(c, h)
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[5].height = 20

    if not activity_stats:
        ws.merge_cells(start_row=6, start_column=second_table_col, end_row=6, end_column=second_table_col + 3)
        c = ws.cell(row=6, column=second_table_col)
        c.value = "Нет данных за выбранный период"
        c.alignment = Alignment(horizontal="center")
        c.font = Font(italic=True, color="999999")
        last_row_table2 = 8
    else:
        for idx, row in enumerate(activity_stats, start=1):
            row_idx = 5 + idx
            dte = row.get("date")
            dte_str = dte.strftime("%d.%m.%Y") if hasattr(dte, "strftime") else (str(dte) if dte else "—")
            type_raw = str(row.get("event_type", "—"))
            type_str = ACTIVITY_MAP.get(type_raw.upper(), type_raw)
            values = [row.get("cow_tag", "—"), dte_str, type_str, row.get("event_count", 0)]
            for col_offset, val in enumerate(values):
                cell = ws.cell(row=row_idx, column=second_table_col + col_offset)
                _data(cell, val, idx)
            ws.row_dimensions[row_idx].height = 18
        last_row_table2 = 5 + len(activity_stats) + 2

    last_row = max(next_row, last_row_table2)

    # ── Подпись ──
    ws.merge_cells(start_row=last_row, start_column=1, end_row=last_row, end_column=6)
    f = ws.cell(row=last_row, column=1)
    f.value = f"Документ сформирован автоматически системой «{company_name}» | {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    f.font = Font(italic=True, size=8, color="999999", name="Calibri")
    f.alignment = Alignment(horizontal="center")

    wb.save(save_path)
    return save_path