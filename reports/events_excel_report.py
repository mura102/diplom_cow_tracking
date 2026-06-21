import os
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    raise ImportError("Отсутствует библиотека openpyxl. Установите её: pip install openpyxl")

# --- Константы стиля (как в старом отчете) ---
COLOR_HEADER = "1F4E79"
COLOR_ROW_ODD = "EBF3FB"
COLOR_ROW_EVEN = "FFFFFF"

COLOR_CRITICAL_FILL = PatternFill("solid", fgColor="F8CECC")  # Красно-розовый для критических
COLOR_WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")  # Желтоватый для предупреждений

THIN_BORDER = Border(
    left=Side(style="thin", color="B0C4DE"),
    right=Side(style="thin", color="B0C4DE"),
    top=Side(style="thin", color="B0C4DE"),
    bottom=Side(style="thin", color="B0C4DE"),
)


def _get_any_field(item, *field_names):
    """Ищет значение по всем возможным названиям ключей/атрибутов."""
    for f in field_names:
        if isinstance(item, dict) and f in item:
            return item[f]
        if hasattr(item, '_mapping') and f in item._mapping:
            return item._mapping[f]
        if hasattr(item, f):
            return getattr(item, f)
    return None


def generate_events_excel_report(events, save_path, company_name, days):
    """
    Генерирует Excel-отчет по уведомлениям, повторяющий оригинальный стиль.
    Колонка "Метка коровы" убрана.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчёт по уведомлениям"
    ws.sheet_view.showGridLines = False

    # --- Заголовок документа ---
    ws.merge_cells('A1:D1')
    title_cell = ws['A1']
    title_cell.value = f"📋 Отчёт по аномалиям — {company_name}"
    title_cell.font = Font(size=14, bold=True, color="1F4E79", name="Calibri")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.fill = PatternFill("solid", fgColor="D6E4F0")
    ws.row_dimensions[1].height = 28

    # --- Дата формирования ---
    ws.merge_cells('A2:D2')
    date_cell = ws['A2']
    date_cell.value = f"Сформирован: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    date_cell.font = Font(italic=True, size=10, color="666666", name="Calibri")
    date_cell.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[2].height = 18

    ws.row_dimensions[3].height = 5  # Отступ

    # --- Заголовки таблицы ---
    headers = ["№", "Сообщение", "Уровень", "Дата и время"]
    header_row = 4

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num, value=header)
        cell.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
        cell.fill = PatternFill("solid", fgColor=COLOR_HEADER)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[header_row].height = 20

    # --- Заполнение данными ---
    row_num = header_row + 1

    if not events:
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=4)
        c = ws.cell(row=row_num, column=1, value="Нет данных за выбранный период")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.font = Font(italic=True, color="999999", name="Calibri")
        c.border = THIN_BORDER
    else:
        for i, event in enumerate(events, 1):
            dt_str = "—"
            severity = "—"
            message = "—"

            try:
                # Дата
                dt_raw = _get_any_field(event, 'triggered_at', 'date_time', 'created_at', 'start_time', 'date',
                                        'timestamp')
                if dt_raw:
                    if hasattr(dt_raw, 'strftime'):
                        dt_str = dt_raw.strftime("%d.%m.%Y %H:%M")
                    else:
                        dt_str = str(dt_raw)

                # Уровень с переводом
                sev_raw = _get_any_field(event, 'severity', 'level', 'alert_level')
                if sev_raw in ("CRITICAL", "Критическое", "Критический"):
                    severity = "Критическое"
                elif sev_raw in ("WARNING", "Предупреждение", "Внимание"):
                    severity = "Предупреждение"
                elif sev_raw:
                    severity = str(sev_raw)

                # Сообщение
                msg_raw = _get_any_field(event, 'message', 'msg', 'description')
                if isinstance(msg_raw, bytes):
                    message = msg_raw.decode("cp1251", errors="replace")
                elif msg_raw:
                    message = str(msg_raw)

            except Exception as e:
                message = f"Ошибка извлечения данных: {str(e)}"

            # Если всё пусто, выводим дамп ключей для дебага
            if dt_str == "—" and severity == "—" and message == "—":
                keys = list(event.keys()) if isinstance(event, dict) else dir(event)
                message = f"[ОШИБКА ДАННЫХ] Доступные ключи в БД: {keys}"

            # Формируем строку (колонка метки удалена)
            row_data = [i, message, severity, dt_str]

            # Чередование цветов строк (Белый / Голубой)
            row_fill = PatternFill("solid", fgColor=COLOR_ROW_EVEN if i % 2 != 0 else COLOR_ROW_ODD)

            for col_num, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num, value=val)
                cell.border = THIN_BORDER
                cell.font = Font(name="Calibri", size=10)
                cell.fill = row_fill

                # Центруем всё, кроме сообщения
                if col_num in (1, 3, 4):  # №, Уровень, Дата
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                else:  # Сообщение
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

                # Цветовая заливка для колонки "Уровень" (теперь она 3-я)
                if col_num == 3:
                    if severity == "Критическое":
                        cell.fill = COLOR_CRITICAL_FILL
                    elif severity == "Предупреждение":
                        cell.fill = COLOR_WARNING_FILL

            ws.row_dimensions[row_num].height = 25  # Увеличиваем высоту для длинных сообщений
            row_num += 1

    # --- Настройка ширины колонок ---
    # Перекинули ширину старой колонки "Метка" (15) в Сообщение (60 -> 75)
    widths = [6, 75, 16, 18]
    for col_num, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_num)].width = width

    wb.save(save_path)